from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document as WordDocument
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from app.chunking import split_markdown
from app.domain import ChunkDraft

SUPPORTED_SUFFIXES = {".md", ".txt", ".pdf", ".docx", ".pptx", ".xlsx"}


def parse_document(filename: str, raw: bytes) -> list[ChunkDraft]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".md", ".txt"}:
        return split_markdown(raw.decode("utf-8"))
    if suffix == ".pdf":
        return _sections_from_pdf(raw)
    if suffix == ".docx":
        return _sections_from_word(raw)
    if suffix == ".pptx":
        return _sections_from_presentation(raw)
    if suffix == ".xlsx":
        return _sections_from_workbook(raw)
    raise ValueError("unsupported document type")


def _sections_from_pdf(raw: bytes) -> list[ChunkDraft]:
    reader = PdfReader(BytesIO(raw))
    return [
        ChunkDraft(heading=f"第 {index} 页", content=page_text)
        for index, page in enumerate(reader.pages, start=1)
        if (page_text := page.extract_text().strip())
    ]


def _sections_from_word(raw: bytes) -> list[ChunkDraft]:
    document = WordDocument(BytesIO(raw))
    sections: list[ChunkDraft] = []
    heading: str | None = None
    buffer: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style.name.startswith("Heading"):
            if buffer:
                sections.append(ChunkDraft(heading=heading, content="\n".join(buffer)))
            heading, buffer = text, []
        else:
            buffer.append(text)
    if buffer:
        sections.append(ChunkDraft(heading=heading, content="\n".join(buffer)))
    return sections


def _sections_from_presentation(raw: bytes) -> list[ChunkDraft]:
    presentation = Presentation(BytesIO(raw))
    sections: list[ChunkDraft] = []
    for index, slide in enumerate(presentation.slides, start=1):
        text = "\n".join(
            shape.text.strip()
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text.strip()
        )
        if text:
            sections.append(ChunkDraft(heading=f"第 {index} 张幻灯片", content=text))
    return sections


def _sections_from_workbook(raw: bytes) -> list[ChunkDraft]:
    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    sections: list[ChunkDraft] = []
    for sheet in workbook.worksheets:
        rows = [
            " | ".join(str(value) for value in row if value is not None)
            for row in sheet.iter_rows(values_only=True)
        ]
        text = "\n".join(row for row in rows if row)
        if text:
            sections.append(ChunkDraft(heading=f"工作表：{sheet.title}", content=text))
    return sections

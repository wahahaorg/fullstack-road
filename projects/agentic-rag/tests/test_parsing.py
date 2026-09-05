from io import BytesIO

from docx import Document as WordDocument
from openpyxl import Workbook
from pptx import Presentation

from app.parsing import parse_document


def test_word_parser_preserves_heading_sections() -> None:
    document = WordDocument()
    document.add_heading("住宿标准", level=1)
    document.add_paragraph("上海每晚 650 元。")
    output = BytesIO()
    document.save(output)

    chunks = parse_document("travel.docx", output.getvalue())

    assert chunks[0].heading == "住宿标准"
    assert "650" in chunks[0].content


def test_presentation_parser_preserves_slide_position() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "故障响应"
    output = BytesIO()
    presentation.save(output)

    chunks = parse_document("oncall.pptx", output.getvalue())

    assert chunks[0].heading == "第 1 张幻灯片"
    assert chunks[0].content == "故障响应"


def test_workbook_parser_preserves_worksheet_name_and_cells() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "住宿标准"
    sheet.append(["城市", "上限"])
    sheet.append(["上海", 650])
    output = BytesIO()
    workbook.save(output)

    chunks = parse_document("travel.xlsx", output.getvalue())

    assert chunks[0].heading == "工作表：住宿标准"
    assert "上海 | 650" in chunks[0].content
